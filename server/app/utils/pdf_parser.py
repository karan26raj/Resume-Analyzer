import pymupdf
from docx import Document


def extract_text_from_pdf(file_path: str) -> str:
    with pymupdf.open(file_path) as document:
        return "\n".join(page.get_text() for page in document).strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract paragraph and table text from a DOCX document."""
    document = Document(file_path)
    text_parts = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            text_parts.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(part for part in text_parts if part.strip()).strip()
